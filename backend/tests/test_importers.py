import io
import json
import zipfile

import docx
import pytest

from smartbook.importers import ImportErrorBadFile, normalize, parse_upload


def make_pdf(text: str, title: str) -> bytes:
    content = b"BT /F1 12 Tf 72 720 Td (" + text.encode() + b") Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Title (" + title.encode() + b") /Author (Jane Doe) >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R /Info 6 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return out


def make_docx() -> bytes:
    document = docx.Document()
    document.core_properties.title = "The Quiet Garden"
    document.core_properties.author = "Mara Lind"
    document.core_properties.keywords = "gardening, patience"
    document.add_paragraph("A gardener spends a year tending a walled garden.")
    document.add_paragraph("Each season teaches something about patience.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_epub() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   "</rootfiles></container>")
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                   '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>The Lighthouse Keeper</dc:title>'
                   "<dc:creator>Ada Stone</dc:creator><dc:subject>Solitude</dc:subject><dc:subject>The sea</dc:subject>"
                   "<dc:description>&lt;p&gt;A keeper alone on a rock.&lt;/p&gt;</dc:description></metadata>"
                   '<manifest><item id="c1" href="text/ch1.xhtml" media-type="application/xhtml+xml"/>'
                   '<item id="c2" href="text/ch2.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
        z.writestr("OEBPS/text/ch1.xhtml", "<html><head><title>skip me</title><style>p{}</style></head>"
                                           "<body><h1>Chapter One</h1><p>The light turned all night.</p></body></html>")
        z.writestr("OEBPS/text/ch2.xhtml", "<html><body><p>Morning came   grey and slow.</p></body></html>")
    return buf.getvalue()


def test_pdf():
    [draft] = parse_upload("keeper.pdf", make_pdf("The lamp burned until dawn.", "Night Watch"))
    assert draft["title"] == "Night Watch" and draft["author"] == "Jane Doe"
    assert "lamp burned until dawn" in draft["text"] and draft["source"] == "pdf"


def test_docx():
    [draft] = parse_upload("garden.docx", make_docx())
    assert (draft["title"], draft["author"]) == ("The Quiet Garden", "Mara Lind")
    assert draft["themes"] == ["gardening", "patience"]
    assert draft["text"].startswith("A gardener spends a year")
    assert draft["description"].startswith("A gardener")  # falls back to the opening text


def test_epub():
    [draft] = parse_upload("keeper.epub", make_epub())
    assert (draft["title"], draft["author"]) == ("The Lighthouse Keeper", "Ada Stone")
    assert draft["themes"] == ["Solitude", "The sea"]
    assert draft["description"] == "A keeper alone on a rock."
    assert draft["text"] == "Chapter One\n\nThe light turned all night.\n\nMorning came grey and slow."


def test_markdown_front_matter():
    md = b"---\ntitle: Salt and Stone\nauthor: R. Vale\ngenres: Fantasy, Adventure\n---\n# Ignored heading\n\nBody text."
    [draft] = parse_upload("notes.md", md)
    assert (draft["title"], draft["author"], draft["genres"]) == ("Salt and Stone", "R. Vale", ["Fantasy", "Adventure"])
    assert draft["text"].startswith("# Ignored heading")


def test_markdown_heading_and_by_line():
    [draft] = parse_upload("x.markdown", b"# The River\n*by Ana Moss*\n\nA long river journey.")
    assert (draft["title"], draft["author"]) == ("The River", "Ana Moss")


def test_txt_uses_filename():
    [draft] = parse_upload("my_reading_notes.txt", b"Some notes about a book.")
    assert draft["title"] == "my reading notes" and draft["author"] == "Unknown"


@pytest.mark.parametrize("payload", [
    [{"title": "A", "author": "X", "summary": "s", "tags": "one; two"}, {"title": "B"}],
    {"books": [{"title": "A", "author": "X", "summary": "s", "tags": ["one", "two"]}, {"title": "B"}]},
])
def test_json_lists(payload):
    drafts = parse_upload("books.json", json.dumps(payload).encode())
    assert [d["title"] for d in drafts] == ["A", "B"]
    assert drafts[0]["themes"] == ["one", "two"] and drafts[0]["description"] == "s"
    assert drafts[1]["description"] == "No description provided."


def test_json_single_object_and_bad_rows():
    [draft] = parse_upload("b.json", json.dumps({"title": "Solo", "genre": "Poetry"}).encode())
    assert draft["genres"] == ["Poetry"]
    [draft] = parse_upload("b.json", json.dumps([{"title": ""}, {"title": "Ok"}, 5]).encode())
    assert draft["title"] == "Ok"


@pytest.mark.parametrize("name, data, match", [
    ("book.exe", b"MZ", "unsupported"),
    ("book.md", b"", "empty"),
    ("book.json", b"{nope", "invalid JSON"),
    ("book.json", b'[{"title": ""}]', "missing its title"),
    ("book.epub", b"not a zip", "unreadable EPUB"),
    ("book.pdf", b"%PDF-garbage", "PDF"),
])
def test_bad_files(name, data, match):
    with pytest.raises(ImportErrorBadFile, match=match):
        parse_upload(name, data)


def test_normalize_limits():
    book = normalize({"title": "T" * 500, "author": "", "genres": [" a ", "", "b"]})
    assert len(book["title"]) == 300 and book["author"] == "Unknown" and book["genres"] == ["a", "b"]
